from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]


def clean_inline(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^\*\*(.+)\*\*$", r"\1", text)
    text = text.replace("**", "")
    text = text.replace("`", "")
    return text


def split_table_row(line: str) -> list[str]:
    return [clean_inline(cell) for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def parse_blocks(markdown: str) -> list[tuple[str, object]]:
    lines = markdown.splitlines()
    blocks: list[tuple[str, object]] = []
    paragraph: list[str] = []
    i = 0

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append(("paragraph", clean_inline(" ".join(paragraph))))
            paragraph = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append((f"heading{min(level, 3)}", clean_inline(stripped[level:].strip())))
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            flush_paragraph()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_separator(lines[i]):
                    table_lines.append(split_table_row(lines[i]))
                i += 1
            blocks.append(("table", table_lines))
            continue

        if stripped == "---":
            flush_paragraph()
            blocks.append(("pagebreak", None))
            i += 1
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            blocks.append(("bullet", clean_inline(stripped[2:])))
            i += 1
            continue

        paragraph.append(stripped)
        i += 1

    flush_paragraph()
    return blocks


def markdown_to_docx(source: Path, target: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    for kind, payload in parse_blocks(source.read_text(encoding="utf-8")):
        if kind.startswith("heading"):
            level = int(kind[-1])
            doc.add_heading(str(payload), level=level)
        elif kind == "paragraph":
            doc.add_paragraph(str(payload))
        elif kind == "bullet":
            doc.add_paragraph(str(payload), style="List Bullet")
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "table":
            rows = payload  # type: ignore[assignment]
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for row_idx, row in enumerate(rows):
                for col_idx, cell_value in enumerate(row):
                    cell = table.cell(row_idx, col_idx)
                    cell.text = cell_value
                    if row_idx == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()

    doc.save(target)


def markdown_to_pdf(source: Path, target: Path) -> None:
    doc = SimpleDocTemplate(
        str(target),
        pagesize=landscape(A4),
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
    )
    styles = getSampleStyleSheet()
    story = []

    for kind, payload in parse_blocks(source.read_text(encoding="utf-8")):
        if kind == "heading1":
            story.append(Paragraph(str(payload), styles["Title"]))
            story.append(Spacer(1, 0.12 * inch))
        elif kind == "heading2":
            story.append(Paragraph(str(payload), styles["Heading2"]))
            story.append(Spacer(1, 0.08 * inch))
        elif kind == "heading3":
            story.append(Paragraph(str(payload), styles["Heading3"]))
            story.append(Spacer(1, 0.06 * inch))
        elif kind == "paragraph":
            story.append(Paragraph(str(payload), styles["BodyText"]))
            story.append(Spacer(1, 0.08 * inch))
        elif kind == "bullet":
            story.append(Paragraph("• " + str(payload), styles["BodyText"]))
            story.append(Spacer(1, 0.04 * inch))
        elif kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "table":
            rows = payload  # type: ignore[assignment]
            if not rows:
                continue
            wrapped = [[Paragraph(str(cell), styles["BodyText"]) for cell in row] for row in rows]
            table = Table(wrapped, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 0.12 * inch))

    doc.build(story)


def main() -> None:
    sources = [
        ROOT / "chapter4_experimental_results_revisions.md",
        ROOT / "submission_inspection_report.md",
    ]
    for source in sources:
        markdown_to_docx(source, source.with_suffix(".docx"))
        markdown_to_pdf(source, source.with_suffix(".pdf"))
        print(f"Generated {source.with_suffix('.docx').name}")
        print(f"Generated {source.with_suffix('.pdf').name}")


if __name__ == "__main__":
    main()
